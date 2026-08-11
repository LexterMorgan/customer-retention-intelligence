#!/usr/bin/env python3
"""Build the pre-native Excel customer retention dashboard.

This script implements Phase A from EXCEL_DASHBOARD_BUILD_PLAN.md. It creates
the complete data, calculation, QA, and presentation scaffold. Native
PivotTables, PivotCharts, slicers, and report connections remain intentionally
reserved for Microsoft Excel Desktop.
"""

from __future__ import annotations

import argparse
import hashlib
import math
import sys
import zipfile
from pathlib import Path
from typing import Any

import pandas as pd
import xlsxwriter
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "processed" / "customers_clean.csv"
DEFAULT_OUTPUT = ROOT / "excel" / "customer_retention_dashboard.xlsx"

TENURE_ORDER = [
    "0-6 months",
    "7-12 months",
    "13-24 months",
    "25-36 months",
    "37-48 months",
    "49-72 months",
]
CONTRACT_ORDER = ["Month-to-Month", "One Year", "Two Year"]
RISK_ORDER = ["Very High", "High", "Medium", "Low"]

EXPECTED = {
    "total": 7043,
    "existing": 6589,
    "churned": 1869,
    "retained": 4720,
    "joined": 454,
    "churn_rate": 0.2837,
    "retention_rate": 0.7163,
    "mrvl": 137086.65,
}

COLORS = {
    "bg": "#F4F6F8",
    "white": "#FFFFFF",
    "navy": "#16324F",
    "blue": "#3E6B89",
    "blue_light": "#DDE9F0",
    "teal": "#2A9D8F",
    "teal_light": "#DDF1EE",
    "coral": "#C94C4C",
    "coral_light": "#F8E3E3",
    "text": "#25313C",
    "muted": "#66727D",
    "border": "#DCE2E7",
    "heat_min": "#FFF9F4",
    "filter_bg": "#EAF0F4",
    "placeholder": "#F8FAFB",
    "success": "#E8F4F1",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def assert_close(actual: float, expected: float, tolerance: float = 0.01) -> None:
    if abs(float(actual) - float(expected)) > tolerance:
        raise AssertionError(f"Expected {expected}, got {actual}")


def load_validate_and_enrich() -> tuple[pd.DataFrame, dict[str, Any]]:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    source_columns = pd.read_csv(SOURCE, nrows=0).columns.tolist()
    if len(source_columns) != 46:
        raise AssertionError(f"Expected 46 source columns, got {len(source_columns)}")

    df = pd.read_csv(SOURCE, keep_default_na=False, na_values=[""])
    if len(df) != EXPECTED["total"]:
        raise AssertionError(f"Expected 7,043 rows, got {len(df):,}")
    if df["Customer ID"].nunique() != EXPECTED["total"]:
        raise AssertionError("Customer ID grain is not one row per customer")
    if df["Customer ID"].duplicated().any():
        raise AssertionError("Duplicate Customer IDs found")
    if set(df["Customer Status"].dropna().unique()) != {"Churned", "Stayed", "Joined"}:
        raise AssertionError("Unexpected Customer Status values")

    existing_mask = df["Customer Status"].isin(["Churned", "Stayed"])
    churned_mask = df["Customer Status"].eq("Churned")
    retained_mask = df["Customer Status"].eq("Stayed")
    joined_mask = df["Customer Status"].eq("Joined")

    metrics = {
        "total": len(df),
        "existing": int(existing_mask.sum()),
        "churned": int(churned_mask.sum()),
        "retained": int(retained_mask.sum()),
        "joined": int(joined_mask.sum()),
        "churn_rate": float(churned_mask.sum() / existing_mask.sum()),
        "retention_rate": float(retained_mask.sum() / existing_mask.sum()),
        "mrvl": round(float(df.loc[churned_mask, "Monthly Charge"].sum()), 2),
    }

    for key in ["total", "existing", "churned", "retained", "joined"]:
        if metrics[key] != EXPECTED[key]:
            raise AssertionError(f"{key}: expected {EXPECTED[key]}, got {metrics[key]}")
    assert_close(metrics["churn_rate"], EXPECTED["churn_rate"], 0.00005)
    assert_close(metrics["retention_rate"], EXPECTED["retention_rate"], 0.00005)
    assert_close(metrics["mrvl"], EXPECTED["mrvl"])

    tenure_sort = {name: idx + 1 for idx, name in enumerate(TENURE_ORDER)}
    tenure_slicer = {
        name: f"{idx + 1:02d} | {name}" for idx, name in enumerate(TENURE_ORDER)
    }

    df["Existing_Flag"] = existing_mask.astype(int)
    df["Churn_Rate_Value"] = df["Is_Churned"].where(existing_mask)
    df["Retention_Rate_Value"] = df["Is_Retained"].where(existing_mask)
    df["MRVL_Row"] = df["Monthly Charge"].where(churned_mask, 0)
    df["Tenure_Band_Sort"] = df["Tenure_Band"].map(tenure_sort)
    df["Tenure_Band_Slicer"] = df["Tenure_Band"].map(tenure_slicer)

    df["Risk_Points"] = (
        df["Contract"].eq("Month-to-Month").astype(int) * 3
        + df["Tenure_Band"].eq("0-6 months").astype(int) * 3
        + df["Internet Type"].eq("Fiber Optic").astype(int) * 2
        + df["Number of Dependents"].eq(0).astype(int)
        + df["Married"].eq("No").astype(int)
    )

    def tier(points: int) -> str:
        if points >= 6:
            return "Very High"
        if points >= 4:
            return "High"
        if points >= 2:
            return "Medium"
        return "Low"

    df["Descriptive_Churn_Risk_Tier"] = df["Risk_Points"].map(tier)
    df["Risk_Tier_Sort"] = df["Descriptive_Churn_Risk_Tier"].map(
        {"Very High": 1, "High": 2, "Medium": 3, "Low": 4}
    )

    if len(df.columns) != 55:
        raise AssertionError(f"Expected 55 workbook columns, got {len(df.columns)}")
    if int(df["Existing_Flag"].sum()) != EXPECTED["existing"]:
        raise AssertionError("Existing_Flag does not reconcile")
    assert_close(df["MRVL_Row"].sum(), EXPECTED["mrvl"])
    if df.loc[joined_mask, "Churn_Rate_Value"].notna().any():
        raise AssertionError("Joined rows must have blank Churn_Rate_Value")

    base = df.loc[existing_mask].copy()

    contract = (
        base.groupby("Contract", observed=True)
        .agg(existing=("Customer ID", "count"), churned=("Is_Churned", "sum"))
    )
    contract["rate"] = contract["churned"] / contract["existing"]
    assert contract.loc["Month-to-Month", "existing"] == 3202
    assert contract.loc["Month-to-Month", "churned"] == 1655
    assert_close(contract.loc["Month-to-Month", "rate"], 0.5169, 0.00005)

    tenure = (
        base.groupby("Tenure_Band", observed=True)
        .agg(existing=("Customer ID", "count"), churned=("Is_Churned", "sum"))
    )
    tenure["rate"] = tenure["churned"] / tenure["existing"]
    assert tenure.loc["0-6 months", "existing"] == 1016
    assert tenure.loc["0-6 months", "churned"] == 784
    assert_close(tenure.loc["0-6 months", "rate"], 0.7717, 0.00005)

    internet = (
        base.loc[base["Internet Service"].eq("Yes")]
        .groupby("Internet Type", observed=True)
        .agg(existing=("Customer ID", "count"), churned=("Is_Churned", "sum"))
    )
    internet["rate"] = internet["churned"] / internet["existing"]
    assert internet.loc["Fiber Optic", "existing"] == 2934
    assert internet.loc["Fiber Optic", "churned"] == 1236
    assert_close(internet.loc["Fiber Optic", "rate"], 0.4213, 0.00005)

    top_reasons = (
        df.loc[churned_mask, "Churn Reason"]
        .value_counts()
        .head(5)
        .rename_axis("reason")
        .reset_index(name="count")
    )
    if top_reasons.iloc[0]["count"] != 313 or top_reasons.iloc[1]["count"] != 311:
        raise AssertionError("Top churn reason benchmarks do not reconcile")

    risk = (
        base.groupby("Descriptive_Churn_Risk_Tier", observed=True)
        .agg(existing=("Customer ID", "count"), churned=("Is_Churned", "sum"))
    )
    risk["rate"] = risk["churned"] / risk["existing"]
    risk_expected = {
        "Very High": (2156, 0.6651),
        "High": (1259, 0.2065),
        "Medium": (1714, 0.0729),
        "Low": (1460, 0.0342),
    }
    for name, (count, rate) in risk_expected.items():
        if risk.loc[name, "existing"] != count:
            raise AssertionError(f"{name} tier count mismatch")
        assert_close(risk.loc[name, "rate"], rate, 0.00005)

    high_risk = base[
        base["Contract"].eq("Month-to-Month")
        & base["Tenure_Band"].eq("0-6 months")
        & base["Internet Type"].eq("Fiber Optic")
    ]
    high_risk_metrics = {
        "existing": len(high_risk),
        "churned": int(high_risk["Is_Churned"].sum()),
        "rate": float(high_risk["Is_Churned"].mean()),
    }
    if high_risk_metrics["existing"] != 487 or high_risk_metrics["churned"] != 444:
        raise AssertionError("High-risk intersection mismatch")
    assert_close(high_risk_metrics["rate"], 0.9117, 0.00005)

    heatmap = (
        base.groupby(["Tenure_Band", "Contract"], observed=True)
        .agg(existing=("Customer ID", "count"), churned=("Is_Churned", "sum"))
        .reset_index()
    )
    heatmap["rate"] = heatmap["churned"] / heatmap["existing"]

    validation = {
        "metrics": metrics,
        "contract": contract,
        "tenure": tenure,
        "internet": internet,
        "top_reasons": top_reasons,
        "risk": risk,
        "high_risk": high_risk_metrics,
        "heatmap": heatmap,
        "source_columns": source_columns,
    }
    return df, validation


def add_outer_border(
    worksheet: xlsxwriter.worksheet.Worksheet,
    first_row: int,
    first_col: int,
    last_row: int,
    last_col: int,
    workbook: xlsxwriter.Workbook,
    fill: str = COLORS["white"],
) -> None:
    """Apply a clean filled panel with only an outer border."""
    for row in range(first_row, last_row + 1):
        for col in range(first_col, last_col + 1):
            props: dict[str, Any] = {"bg_color": fill}
            if row == first_row:
                props["top"] = 1
                props["top_color"] = COLORS["border"]
            if row == last_row:
                props["bottom"] = 1
                props["bottom_color"] = COLORS["border"]
            if col == first_col:
                props["left"] = 1
                props["left_color"] = COLORS["border"]
            if col == last_col:
                props["right"] = 1
                props["right_color"] = COLORS["border"]
            worksheet.write_blank(row, col, None, workbook.add_format(props))


def write_formula_merged(
    worksheet: xlsxwriter.worksheet.Worksheet,
    cell_range: str,
    formula: str,
    fmt: xlsxwriter.format.Format,
    cached_value: Any,
) -> None:
    first_cell = cell_range.split(":")[0]
    worksheet.merge_range(cell_range, "", fmt)
    worksheet.write_formula(first_cell, formula, fmt, cached_value)


def blend_heat_color(value: float) -> str:
    value = max(0.0, min(1.0, value))
    start = (255, 249, 244)
    end = (201, 76, 76)
    rgb = tuple(round(start[i] + (end[i] - start[i]) * value) for i in range(3))
    return "#{:02X}{:02X}{:02X}".format(*rgb)


def build_workbook(
    output: Path, df: pd.DataFrame, validation: dict[str, Any], source_hash: str
) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        output.unlink()

    workbook = xlsxwriter.Workbook(
        output,
        {
            "nan_inf_to_errors": True,
            "use_zip64": False,
        },
    )
    workbook.set_properties(
        {
            "title": "Customer Retention Intelligence Dashboard",
            "subject": "Interactive Excel churn and retention analytics",
            "author": "Customer Retention Intelligence",
            "company": "Portfolio Analytics",
            "comments": (
                "Phase A scaffold. Native PivotTables, PivotCharts, and slicers "
                "must be added in Microsoft Excel Desktop."
            ),
        }
    )
    workbook.set_calc_mode("auto")

    dashboard = workbook.add_worksheet("01_Dashboard")
    pivots = workbook.add_worksheet("02_Pivots")
    calc = workbook.add_worksheet("03_Calculations")
    data_ws = workbook.add_worksheet("04_Data")

    dashboard.set_tab_color(COLORS["navy"])
    pivots.set_tab_color(COLORS["blue"])
    calc.set_tab_color(COLORS["teal"])
    data_ws.set_tab_color(COLORS["muted"])

    # Shared support formats.
    title_fmt = workbook.add_format(
        {
            "font_name": "Aptos Display",
            "font_size": 20,
            "bold": True,
            "font_color": COLORS["white"],
            "bg_color": COLORS["navy"],
            "valign": "vcenter",
        }
    )
    header_subtitle_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 10,
            "font_color": "#D7E3EB",
            "bg_color": COLORS["navy"],
            "valign": "vcenter",
        }
    )
    badge_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 9,
            "bold": True,
            "font_color": COLORS["white"],
            "bg_color": COLORS["blue"],
            "align": "center",
            "valign": "vcenter",
            "border": 0,
        }
    )
    section_title_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 12,
            "bold": True,
            "font_color": COLORS["navy"],
            "bg_color": COLORS["white"],
            "valign": "vcenter",
        }
    )
    section_subtitle_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 8,
            "font_color": COLORS["muted"],
            "bg_color": COLORS["white"],
            "valign": "vcenter",
        }
    )
    note_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 8,
            "italic": True,
            "font_color": COLORS["muted"],
            "bg_color": COLORS["white"],
            "valign": "vcenter",
        }
    )
    placeholder_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 8,
            "font_color": "#8A98A4",
            "bg_color": COLORS["placeholder"],
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "border_color": "#EEF2F5",
        }
    )
    support_title_fmt = workbook.add_format(
        {
            "font_name": "Aptos Display",
            "font_size": 18,
            "bold": True,
            "font_color": COLORS["navy"],
        }
    )
    support_header_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 9,
            "bold": True,
            "font_color": COLORS["white"],
            "bg_color": COLORS["navy"],
            "border": 1,
            "border_color": COLORS["navy"],
            "align": "center",
            "valign": "vcenter",
        }
    )
    support_cell_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 9,
            "font_color": COLORS["text"],
            "border": 1,
            "border_color": COLORS["border"],
            "valign": "top",
            "text_wrap": True,
        }
    )
    pass_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 9,
            "bold": True,
            "font_color": COLORS["teal"],
            "bg_color": COLORS["success"],
            "align": "center",
        }
    )
    count_fmt = workbook.add_format({"num_format": "#,##0", "font_name": "Aptos"})
    pct_qa_fmt = workbook.add_format({"num_format": "0.00%", "font_name": "Aptos"})
    money_qa_fmt = workbook.add_format(
        {"num_format": "$#,##0.00", "font_name": "Aptos"}
    )

    # 04_Data.
    data_ws.hide_gridlines(2)
    data_ws.freeze_panes(1, 0)
    data_ws.set_zoom(80)
    data_ws.set_default_row(15)
    data_ws.set_column(0, len(df.columns) - 1, 13)
    width_map = {
        "Customer ID": 15,
        "City": 20,
        "Zip Code": 10,
        "Offer": 13,
        "Internet Type": 14,
        "Contract": 18,
        "Payment Method": 18,
        "Churn Category": 18,
        "Churn Reason": 38,
        "Tenure_Band": 17,
        "Tenure_Band_Slicer": 22,
        "Descriptive_Churn_Risk_Tier": 28,
    }
    for idx, col in enumerate(df.columns):
        data_ws.set_column(idx, idx, width_map.get(col, 13))

    table_columns = [{"header": str(col)} for col in df.columns]
    data_ws.write_row(0, 0, list(df.columns))
    for row_idx, row in enumerate(df.itertuples(index=False, name=None), start=1):
        cleaned = [None if pd.isna(value) else value for value in row]
        data_ws.write_row(row_idx, 0, cleaned)
    data_ws.add_table(
        0,
        0,
        len(df),
        len(df.columns) - 1,
        {
            "name": "tblCustomers",
            "style": "Table Style Medium 2",
            "columns": table_columns,
            "banded_rows": True,
            "banded_columns": False,
        },
    )

    # 02_Pivots scaffold. Pivot destination cells remain empty.
    pivots.hide_gridlines(2)
    pivots.set_zoom(85)
    pivots.set_column("A:H", 15)
    pivots.set_column("J:J", 23)
    pivots.set_column("K:K", 15)
    pivots.set_column("L:M", 20)
    pivots.set_column("N:Q", 17)
    pivots.write("A1", "Native Pivot Support", support_title_fmt)
    pivots.write(
        "A2",
        "Create all six pivots by copying PT_KPI_Context so every slicer can share one PivotCache.",
        workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 10,
                "font_color": COLORS["coral"],
                "bold": True,
            }
        ),
    )
    inventory = [
        [
            "PivotTable",
            "Destination",
            "Rows",
            "Columns",
            "Values",
            "Filters",
            "Dashboard object",
        ],
        [
            "PT_KPI_Context",
            "A3",
            "—",
            "—",
            "Sum Existing, Churned, Retained, MRVL",
            "—",
            "Six KPI cards",
        ],
        [
            "PT_InternetType",
            "A12",
            "Internet Type",
            "—",
            "Average Churn_Rate_Value",
            "Internet Service = Yes",
            "PC_InternetType",
        ],
        [
            "PT_ChurnReasons",
            "A24",
            "Churn Reason",
            "—",
            "Count Customer ID",
            "Status = Churned; Top 5",
            "PC_ChurnReasons",
        ],
        [
            "PT_RiskTier",
            "A38",
            "Risk Tier",
            "—",
            "Average Churn_Rate_Value",
            "—",
            "PC_RiskTier",
        ],
        [
            "PT_RiskTierVolume",
            "F38",
            "Risk Tier",
            "—",
            "Sum Existing_Flag",
            "—",
            "Risk volume labels",
        ],
        [
            "PT_ContractTenure",
            "A52",
            "Tenure_Band",
            "Contract",
            "Average Churn Rate; Sum Existing",
            "—",
            "Heatmap",
        ],
    ]
    start_row, start_col = 3, 9  # J4
    for r, values in enumerate(inventory):
        for c, value in enumerate(values):
            pivots.write(
                start_row + r,
                start_col + c,
                value,
                support_header_fmt if r == 0 else support_cell_fmt,
            )
    pivots.write(
        "J13",
        "Native Excel boundary",
        workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 11,
                "bold": True,
                "font_color": COLORS["navy"],
            }
        ),
    )
    pivots.merge_range(
        "J14:Q18",
        (
            "Python intentionally does not create PivotTable, PivotChart, slicer, "
            "or report-connection XML. Use the exact Phase B runbook in "
            "EXCEL_DASHBOARD_BUILD_PLAN.md. Do not resave this workbook with "
            "openpyxl or XlsxWriter after native objects are added."
        ),
        workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 9,
                "font_color": COLORS["text"],
                "bg_color": COLORS["filter_bg"],
                "border": 1,
                "border_color": COLORS["border"],
                "text_wrap": True,
                "valign": "top",
            }
        ),
    )

    # 03_Calculations.
    calc.hide_gridlines(2)
    calc.set_zoom(90)
    calc.set_column("A:A", 3)
    calc.set_column("B:B", 28)
    calc.set_column("C:E", 18)
    calc.set_column("F:F", 12)
    calc.set_column("H:H", 20)
    calc.set_column("I:L", 28)
    calc.write("B1", "Calculation & Validation Layer", support_title_fmt)
    calc.write(
        "B2",
        "Current values become slicer-responsive when PT_KPI_Context is created in Excel Desktop.",
        workbook.add_format(
            {"font_name": "Aptos", "font_size": 9, "font_color": COLORS["muted"]}
        ),
    )

    headers = ["KPI", "Current", "Baseline", "Variance", "Status"]
    calc.write_row("B4", headers, support_header_fmt)
    kpi_rows = [
        ("Existing Customers", 6589, "#,##0"),
        ("Churned Customers", 1869, "#,##0"),
        ("Retained Customers", 4720, "#,##0"),
        ("Churn Rate", metrics := validation["metrics"]["churn_rate"], "0.00%"),
        (
            "Retention Rate",
            validation["metrics"]["retention_rate"],
            "0.00%",
        ),
        ("Monthly Recurring Value Lost", 137086.65, "$#,##0.00"),
    ]
    getpivot_formulas = [
        '=IFERROR(GETPIVOTDATA("Existing Customers",\'02_Pivots\'!$A$3),6589)',
        '=IFERROR(GETPIVOTDATA("Churned Customers",\'02_Pivots\'!$A$3),1869)',
        '=IFERROR(GETPIVOTDATA("Retained Customers",\'02_Pivots\'!$A$3),4720)',
        "=IF(C5=0,NA(),C6/C5)",
        "=IF(C5=0,NA(),C7/C5)",
        '=IFERROR(GETPIVOTDATA("Monthly Recurring Value Lost",\'02_Pivots\'!$A$3),137086.65)',
    ]
    for idx, ((label, baseline, number_format), formula) in enumerate(
        zip(kpi_rows, getpivot_formulas), start=5
    ):
        row = idx - 1
        cell_fmt = workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 9,
                "border": 1,
                "border_color": COLORS["border"],
                "num_format": number_format,
            }
        )
        calc.write(row, 1, label, support_cell_fmt)
        calc.write_formula(row, 2, formula, cell_fmt, baseline)
        calc.write(row, 3, baseline, cell_fmt)
        calc.write_formula(row, 4, f"=C{idx}-D{idx}", cell_fmt, 0)
        calc.write_formula(
            row,
            5,
            f'=IF(ABS(E{idx})<0.0001,"PASS","FAIL")',
            pass_fmt,
            "PASS",
        )

    calc.write("H1", "Native pivots ready?", support_header_fmt)
    calc.write_formula(
        "H2",
        '=NOT(ISERROR(GETPIVOTDATA("Existing Customers",\'02_Pivots\'!$A$3)))',
        support_cell_fmt,
        False,
    )
    workbook.define_name("nativePivotsReady", "='03_Calculations'!$H$2")
    name_rows = {
        "kpiExisting": 5,
        "kpiChurned": 6,
        "kpiRetained": 7,
        "kpiChurnRate": 8,
        "kpiRetentionRate": 9,
        "kpiMRVL": 10,
    }
    for name, row in name_rows.items():
        workbook.define_name(name, f"='03_Calculations'!$C${row}")

    qa_rows = [
        ("Total Customers", "=ROWS(tblCustomers[Customer ID])", 7043, "#,##0"),
        (
            "Joined Customers",
            '=COUNTIF(tblCustomers[Customer Status],"Joined")',
            454,
            "#,##0",
        ),
        (
            "Existing identity",
            "=SUM(tblCustomers[Existing_Flag])",
            6589,
            "#,##0",
        ),
        (
            "Rate identity",
            "=SUM(tblCustomers[Is_Churned])/SUM(tblCustomers[Existing_Flag])"
            "+SUM(tblCustomers[Is_Retained])/SUM(tblCustomers[Existing_Flag])",
            1.0,
            "0.00%",
        ),
        (
            "High-risk intersection existing",
            '=COUNTIFS(tblCustomers[Contract],"Month-to-Month",'
            'tblCustomers[Tenure_Band],"0-6 months",'
            'tblCustomers[Internet Type],"Fiber Optic",'
            "tblCustomers[Existing_Flag],1)",
            487,
            "#,##0",
        ),
        (
            "High-risk intersection churned",
            '=COUNTIFS(tblCustomers[Contract],"Month-to-Month",'
            'tblCustomers[Tenure_Band],"0-6 months",'
            'tblCustomers[Internet Type],"Fiber Optic",'
            "tblCustomers[Is_Churned],1)",
            444,
            "#,##0",
        ),
    ]
    calc.write("B13", "Source & denominator QA", section_title_fmt)
    calc.write_row("B14", ["Check", "Formula result", "Expected", "Status"], support_header_fmt)
    for offset, (label, formula, expected, num_fmt) in enumerate(qa_rows, start=15):
        fmt = workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 9,
                "border": 1,
                "border_color": COLORS["border"],
                "num_format": num_fmt,
            }
        )
        calc.write(offset - 1, 1, label, support_cell_fmt)
        calc.write_formula(offset - 1, 2, formula, fmt, expected)
        calc.write(offset - 1, 3, expected, fmt)
        calc.write_formula(
            offset - 1,
            4,
            f'=IF(ABS(C{offset}-D{offset})<0.0001,"PASS","FAIL")',
            pass_fmt,
            "PASS",
        )

    helper_docs = [
        ["Helper field", "Equivalent Excel logic"],
        [
            "Existing_Flag",
            '=--OR([@[Customer Status]]="Churned",[@[Customer Status]]="Stayed")',
        ],
        ["Churn_Rate_Value", '=IF([@Existing_Flag]=1,[@Is_Churned],"")'],
        ["Retention_Rate_Value", '=IF([@Existing_Flag]=1,[@Is_Retained],"")'],
        ["MRVL_Row", '=IF([@Is_Churned]=1,[@[Monthly Charge]],0)'],
        [
            "Risk_Points",
            "M2M +3; 0-6 months +3; Fiber +2; no dependents +1; not married +1",
        ],
        ["Risk Tier", "Very High >=6; High >=4; Medium >=2; otherwise Low"],
    ]
    calc.write("H4", "Helper-field audit", section_title_fmt)
    for r, values in enumerate(helper_docs, start=5):
        row_format = support_header_fmt if r == 5 else support_cell_fmt
        for c, value in enumerate(values, start=7):
            # Formula examples are documentation text. Writing them with
            # worksheet.write() makes XlsxWriter serialize invalid formulas
            # such as [@Existing_Flag] outside an Excel Table, which causes
            # Excel to repair sheet3.xml. Force literal strings permanently.
            calc.write_string(r - 1, c, str(value), row_format)

    checklist = [
        "1. Create PT_KPI_Context at 02_Pivots!A3 from tblCustomers.",
        "2. Copy it to A12, A24, A38, F38, and A52.",
        "3. Configure fields exactly as listed on 02_Pivots.",
        "4. Add three PivotCharts and place them in the prepared wells.",
        "5. Add Contract, Tenure_Band_Slicer, and Internet Type slicers.",
        "6. Connect each slicer to all six PivotTables.",
        "7. Refresh All, recalculate, validate, then hide support sheets.",
    ]
    calc.write("H14", "Native Excel checkpoint", section_title_fmt)
    for r, item in enumerate(checklist, start=15):
        calc.merge_range(r - 1, 7, r - 1, 11, item, support_cell_fmt)

    insight_texts = [
        (
            "insightCritical",
            "Month-to-Month customers have a 51.7% churn rate, the highest contract-level exposure.",
        ),
        (
            "insightDriver",
            "Churn reaches 77.2% among customers in their first 0–6 months.",
        ),
        (
            "insightWarning",
            "Month-to-Month + 0–6 months + Fiber records 91.2% churn: 444 of 487 existing customers.",
        ),
        (
            "insightOpportunity",
            "Fiber churn is 42.1%; better devices (313) and better offers (311) lead stated exit reasons.",
        ),
    ]
    calc.write_row("H24", ["Named insight", "Baseline executive context"], support_header_fmt)
    for row, (name, text) in enumerate(insight_texts, start=25):
        calc.write_string(row - 1, 7, name, support_cell_fmt)
        calc.write_string(row - 1, 8, text, support_cell_fmt)
        workbook.define_name(name, f"='03_Calculations'!$I${row}")

    # Dashboard canvas: consulting-style hierarchy from header to methodology.
    dashboard.hide_gridlines(2)
    dashboard.hide_row_col_headers()
    dashboard.set_zoom(75)
    dashboard.set_landscape()
    dashboard.set_paper(9)
    dashboard.fit_to_pages(1, 1)
    dashboard.set_margins(0.12, 0.12, 0.18, 0.18)
    dashboard.print_area("A1:X58")
    dashboard.activate()
    dashboard.select()

    gutter_cols = {3, 7, 11, 15, 19, 23}
    for col in range(24):
        dashboard.set_column(col, col, 1.4 if col in gutter_cols else 7.3)

    row_heights = {
        1: 27, 2: 19, 3: 18,
        4: 18, 5: 19, 6: 19, 7: 7,
        8: 15, 9: 24, 10: 21, 11: 14, 12: 8,
        13: 22, 14: 22, 15: 20, 28: 8,
        29: 22, 30: 16, 38: 18, 39: 8,
        40: 22, 41: 20, 42: 34, 43: 18, 44: 8,
        45: 22, 46: 19, 47: 19, 48: 19, 49: 19,
        50: 19, 51: 19, 52: 19, 53: 20, 54: 8,
        55: 20, 56: 18, 57: 18, 58: 18,
    }
    for row in range(1, 59):
        dashboard.set_row(row - 1, row_heights.get(row, 16))

    bg_fmt = workbook.add_format({"bg_color": COLORS["bg"]})
    for row in range(58):
        for col in range(24):
            dashboard.write_blank(row, col, None, bg_fmt)

    # Executive header.
    header_fill = workbook.add_format({"bg_color": COLORS["navy"]})
    for row in range(3):
        for col in range(24):
            dashboard.write_blank(row, col, None, header_fill)
    dashboard.merge_range("A1:P2", "CUSTOMER RETENTION INTELLIGENCE", title_fmt)
    dashboard.merge_range(
        "A3:P3",
        "Executive churn health • concentration • drivers • retention opportunity",
        header_subtitle_fmt,
    )
    dashboard.merge_range("Q1:X2", "Q2 2022  •  CALIFORNIA", badge_fmt)
    dashboard.merge_range(
        "Q3:X3",
        "7,043 customers  |  454 joined  |  rates exclude Joined",
        workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 8,
                "font_color": "#D7E3EB",
                "bg_color": COLORS["navy"],
                "align": "center",
                "valign": "vcenter",
            }
        ),
    )

    # Integrated filter bar; the VBA slicers occupy these exact wells.
    filter_band_fmt = workbook.add_format({"bg_color": COLORS["filter_bg"]})
    for row in range(3, 6):
        for col in range(24):
            dashboard.write_blank(row, col, None, filter_band_fmt)
    filter_wells = [
        ("A4:H6", "CONTRACT", "Month-to-Month  •  One Year  •  Two Year"),
        ("I4:P6", "TENURE BAND", "0–6 months  →  49–72 months"),
        (
            "Q4:X6",
            "INTERNET TYPE",
            "Fiber Optic  •  Cable  •  DSL  •  (blank = no internet)",
        ),
    ]
    for cell_range, label, items in filter_wells:
        first, end = cell_range.split(":")
        start_col = ord(first[0]) - ord("A")
        start_row = int(first[1:]) - 1
        end_col = ord(end[0]) - ord("A")
        end_row = int(end[1:]) - 1
        dashboard.merge_range(
            start_row,
            start_col,
            start_row,
            end_col,
            label,
            workbook.add_format(
                {
                    "bg_color": COLORS["white"],
                    "border": 1,
                    "border_color": COLORS["border"],
                    "font_name": "Aptos",
                    "font_size": 8,
                    "bold": True,
                    "font_color": COLORS["navy"],
                    "align": "left",
                    "valign": "vcenter",
                }
            ),
        )
        dashboard.merge_range(
            start_row + 1,
            start_col,
            end_row,
            end_col,
            f"{items}\nNative slicer is created here by setup_dashboard.bas",
            workbook.add_format(
                {
                    "bg_color": COLORS["white"],
                    "border": 1,
                    "border_color": COLORS["border"],
                    "font_name": "Aptos",
                    "font_size": 8,
                    "font_color": COLORS["muted"],
                    "align": "center",
                    "valign": "vcenter",
                    "text_wrap": True,
                }
            ),
        )

    # Six compact executive KPI cards.
    cards = [
        ("A8:C11", "EXISTING CUSTOMERS", "=kpiExisting", 6589, "#,##0", COLORS["blue"], "Active base"),
        ("E8:G11", "CHURNED CUSTOMERS", "=kpiChurned", 1869, "#,##0", COLORS["coral"], "Exited"),
        ("I8:K11", "RETAINED CUSTOMERS", "=kpiRetained", 4720, "#,##0", COLORS["teal"], "Stayed"),
        ("M8:O11", "CHURN RATE", "=kpiChurnRate", validation["metrics"]["churn_rate"], "0.0%", COLORS["coral"], "Churned ÷ Existing"),
        ("Q8:S11", "RETENTION RATE", "=kpiRetentionRate", validation["metrics"]["retention_rate"], "0.0%", COLORS["teal"], "Retained ÷ Existing"),
        ("U8:W11", "MRVL", "=kpiMRVL", 137086.65, "$#,##0", COLORS["coral"], "Monthly value lost"),
    ]
    for cell_range, label, formula, cached, num_format, accent, hint in cards:
        start, end = cell_range.split(":")
        start_col = ord(start[0]) - ord("A")
        start_row = int(start[1:]) - 1
        end_col = ord(end[0]) - ord("A")
        end_row = int(end[1:]) - 1
        for row in range(start_row, end_row + 1):
            for col in range(start_col, end_col + 1):
                props = {"bg_color": COLORS["white"]}
                if row == start_row:
                    props.update({"top": 3, "top_color": accent})
                if row == end_row:
                    props.update({"bottom": 1, "bottom_color": COLORS["border"]})
                if col == start_col:
                    props.update({"left": 1, "left_color": COLORS["border"]})
                if col == end_col:
                    props.update({"right": 1, "right_color": COLORS["border"]})
                dashboard.write_blank(row, col, None, workbook.add_format(props))
        dashboard.merge_range(
            start_row,
            start_col,
            start_row,
            end_col,
            label,
            workbook.add_format(
                {
                    "font_name": "Aptos",
                    "font_size": 8,
                    "bold": True,
                    "font_color": COLORS["muted"],
                    "bg_color": COLORS["white"],
                    "top": 3,
                    "top_color": accent,
                    "left": 1,
                    "left_color": COLORS["border"],
                    "right": 1,
                    "right_color": COLORS["border"],
                    "align": "center",
                    "valign": "vcenter",
                }
            ),
        )
        value_range = f"{chr(65 + start_col)}{start_row + 2}:{chr(65 + end_col)}{start_row + 3}"
        value_fmt = workbook.add_format(
            {
                "font_name": "Aptos Display",
                "font_size": 20,
                "bold": True,
                "font_color": accent,
                "bg_color": COLORS["white"],
                "num_format": num_format,
                "align": "center",
                "valign": "vcenter",
                "left": 1,
                "left_color": COLORS["border"],
                "right": 1,
                "right_color": COLORS["border"],
            }
        )
        write_formula_merged(dashboard, value_range, formula, value_fmt, cached)
        dashboard.merge_range(
            f"{chr(65 + start_col)}{end_row + 1}:{chr(65 + end_col)}{end_row + 1}",
            hint,
            workbook.add_format(
                {
                    "font_name": "Aptos",
                    "font_size": 7,
                    "font_color": COLORS["muted"],
                    "bg_color": COLORS["white"],
                    "bottom": 1,
                    "bottom_color": COLORS["border"],
                    "left": 1,
                    "left_color": COLORS["border"],
                    "right": 1,
                    "right_color": COLORS["border"],
                    "align": "center",
                    "valign": "vcenter",
                }
            ),
        )

    section_band_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 10,
            "bold": True,
            "font_color": COLORS["navy"],
            "bg_color": COLORS["bg"],
            "bottom": 2,
            "bottom_color": COLORS["navy"],
            "valign": "vcenter",
        }
    )
    dashboard.merge_range(
        "A13:X13",
        "MAIN ANALYSIS  |  Concentration • associated characteristics • stated reasons • retention priority",
        section_band_fmt,
    )

    # Four approved analysis panels.
    add_outer_border(dashboard, 13, 0, 26, 15, workbook)
    add_outer_border(dashboard, 13, 16, 26, 23, workbook)
    add_outer_border(dashboard, 28, 0, 37, 12, workbook)
    add_outer_border(dashboard, 28, 13, 37, 23, workbook)
    dashboard.merge_range("A14:P14", "CONTRACT × TENURE CHURN CONCENTRATION", section_title_fmt)
    dashboard.merge_range("Q14:X14", "CHURN RATE BY INTERNET TYPE", section_title_fmt)
    dashboard.merge_range("Q15:X15", "Which internet technology is associated with elevated churn?", section_subtitle_fmt)
    dashboard.merge_range("A29:M29", "TOP 5 STATED CHURN REASONS", section_title_fmt)
    dashboard.merge_range("A30:M30", "What churned customers most often said when leaving", section_subtitle_fmt)
    dashboard.merge_range("N29:X29", "DESCRIPTIVE CHURN RISK TIERS", section_title_fmt)
    dashboard.merge_range("N30:X30", "Where elevated churn and meaningful customer volume overlap", section_subtitle_fmt)

    # Heatmap with baseline cache and slicer-responsive GETPIVOTDATA formulas.
    contract_ranges = {
        "Month-to-Month": "E15:H15",
        "One Year": "I15:L15",
        "Two Year": "M15:P15",
    }
    heat_header_fmt = workbook.add_format(
        {
            "font_name": "Aptos",
            "font_size": 8,
            "bold": True,
            "font_color": COLORS["white"],
            "bg_color": COLORS["blue"],
            "align": "center",
            "valign": "vcenter",
            "border": 1,
            "border_color": COLORS["white"],
        }
    )
    for contract, cell_range in contract_ranges.items():
        dashboard.merge_range(cell_range, contract, heat_header_fmt)
    heatmap = validation["heatmap"].set_index(["Tenure_Band", "Contract"])
    contract_cols = {
        "Month-to-Month": ("E", "H"),
        "One Year": ("I", "L"),
        "Two Year": ("M", "P"),
    }
    for idx, tenure in enumerate(TENURE_ORDER):
        rate_row = 16 + idx * 2
        vol_row = rate_row + 1
        dashboard.merge_range(
            f"A{rate_row}:D{vol_row}",
            tenure,
            workbook.add_format(
                {
                    "font_name": "Aptos",
                    "font_size": 8,
                    "bold": True,
                    "font_color": COLORS["text"],
                    "bg_color": "#F7F9FA",
                    "align": "center",
                    "valign": "vcenter",
                    "border": 1,
                    "border_color": COLORS["white"],
                }
            ),
        )
        for contract, (start_col, end_col) in contract_cols.items():
            values = heatmap.loc[(tenure, contract)]
            rate = float(values["rate"])
            volume = int(values["existing"])
            rate_range = f"{start_col}{rate_row}:{end_col}{rate_row}"
            volume_range = f"{start_col}{vol_row}:{end_col}{vol_row}"
            rate_formula = (
                '=IF(nativePivotsReady,IFERROR(GETPIVOTDATA("Churn Rate",'
                f'\'02_Pivots\'!$A$52,"Tenure_Band","{tenure}",'
                f'"Contract","{contract}"),""),{rate})'
            )
            volume_formula = (
                '=IF(nativePivotsReady,IFERROR(GETPIVOTDATA("Existing Customers",'
                f'\'02_Pivots\'!$A$52,"Tenure_Band","{tenure}",'
                f'"Contract","{contract}"),""),{volume})'
            )
            rate_fmt = workbook.add_format(
                {
                    "font_name": "Aptos Display",
                    "font_size": 11,
                    "bold": True,
                    "font_color": COLORS["text"],
                    "bg_color": blend_heat_color(rate),
                    "num_format": "0.0%",
                    "align": "center",
                    "valign": "vcenter",
                    "border": 1,
                    "border_color": COLORS["white"],
                }
            )
            volume_fmt = workbook.add_format(
                {
                    "font_name": "Aptos",
                    "font_size": 7,
                    "font_color": COLORS["muted"],
                    "bg_color": COLORS["white"],
                    "num_format": '"n="#,##0',
                    "align": "center",
                    "valign": "vcenter",
                    "border": 1,
                    "border_color": COLORS["white"],
                }
            )
            write_formula_merged(dashboard, rate_range, rate_formula, rate_fmt, rate)
            write_formula_merged(dashboard, volume_range, volume_formula, volume_fmt, volume)
            dashboard.conditional_format(
                rate_range,
                {
                    "type": "2_color_scale",
                    "min_type": "num",
                    "min_value": 0,
                    "min_color": COLORS["heat_min"],
                    "max_type": "num",
                    "max_value": 1,
                    "max_color": COLORS["coral"],
                },
            )

    # Polished static previews remain behind the future native PivotCharts.
    preview_label_fmt = workbook.add_format(
        {"font_name": "Aptos", "font_size": 8, "font_color": COLORS["text"], "bg_color": COLORS["placeholder"], "valign": "vcenter"}
    )
    preview_pct_fmt = workbook.add_format(
        {"font_name": "Aptos", "font_size": 8, "bold": True, "font_color": COLORS["text"], "bg_color": COLORS["placeholder"], "num_format": "0.0%", "align": "right", "valign": "vcenter"}
    )
    preview_count_fmt = workbook.add_format(
        {"font_name": "Aptos", "font_size": 8, "bold": True, "font_color": COLORS["text"], "bg_color": COLORS["placeholder"], "num_format": "#,##0", "align": "right", "valign": "vcenter"}
    )
    bar_bg_fmt = workbook.add_format({"bg_color": "#E8EDF1"})
    bar_blue_fmt = workbook.add_format({"bg_color": COLORS["blue"]})
    bar_coral_fmt = workbook.add_format({"bg_color": COLORS["coral"]})
    baseline_tag_fmt = workbook.add_format(
        {"font_name": "Aptos", "font_size": 7, "font_color": COLORS["muted"], "bg_color": COLORS["placeholder"], "align": "center", "valign": "vcenter"}
    )

    dashboard.merge_range("Q16:X16", "BASELINE PREVIEW • NATIVE INTERACTIVITY ADDED BY VBA", baseline_tag_fmt)
    internet_preview = validation["internet"].sort_values("rate", ascending=False)
    for row_num, (name, values) in zip([18, 21, 24], internet_preview.iterrows()):
        rate = float(values["rate"])
        dashboard.merge_range(f"Q{row_num}:S{row_num}", name, preview_label_fmt)
        filled = max(1, min(3, math.ceil(rate / 0.5 * 3)))
        for offset, col in enumerate(range(19, 22)):
            dashboard.write_blank(row_num - 1, col, None, bar_blue_fmt if offset < filled else bar_bg_fmt)
        dashboard.write(row_num - 1, 22, rate, preview_pct_fmt)
    dashboard.merge_range("Q27:X27", "Internet customers only • association, not causation", note_fmt)

    dashboard.merge_range("A31:M31", "BASELINE PREVIEW • NATIVE INTERACTIVITY ADDED BY VBA", baseline_tag_fmt)
    for row_num, reason_row in zip([32, 33, 34, 35, 36], validation["top_reasons"].itertuples()):
        dashboard.merge_range(f"A{row_num}:F{row_num}", reason_row.reason, preview_label_fmt)
        filled = max(1, min(5, math.ceil(int(reason_row.count) / 350 * 5)))
        for offset, col in enumerate(range(6, 11)):
            dashboard.write_blank(row_num - 1, col, None, bar_coral_fmt if offset < filled else bar_bg_fmt)
        dashboard.write(row_num - 1, 12, int(reason_row.count), preview_count_fmt)
    dashboard.merge_range("A38:M38", "Self-reported exit reasons • not verified causes", note_fmt)

    dashboard.merge_range("N31:U31", "BASELINE PREVIEW • NATIVE INTERACTIVITY ADDED BY VBA", baseline_tag_fmt)
    risk = validation["risk"]
    for row_num, tier_name in zip([32, 33, 34, 35], RISK_ORDER):
        rate = float(risk.loc[tier_name, "rate"])
        dashboard.merge_range(f"N{row_num}:P{row_num}", tier_name, preview_label_fmt)
        filled = max(1, min(4, math.ceil(rate / 0.75 * 4)))
        for offset, col in enumerate(range(16, 20)):
            dashboard.write_blank(row_num - 1, col, None, bar_coral_fmt if offset < filled else bar_bg_fmt)
        dashboard.write(row_num - 1, 20, rate, preview_pct_fmt)
    dashboard.merge_range("N38:X38", "Rule-based descriptive tiers • not predictive probabilities", note_fmt)

    # Dynamic risk-tier volumes.
    risk_volume_fmt = workbook.add_format(
        {"font_name": "Aptos", "font_size": 8, "bold": True, "font_color": COLORS["navy"], "bg_color": COLORS["white"], "num_format": "#,##0", "align": "center", "valign": "vcenter", "border": 1, "border_color": COLORS["border"]}
    )
    for row_num, tier_name, baseline in [
        (32, "Very High", 2156), (33, "High", 1259), (34, "Medium", 1714), (35, "Low", 1460)
    ]:
        formula = (
            '=IF(nativePivotsReady,IFERROR(GETPIVOTDATA("Existing Customers",'
            f'\'02_Pivots\'!$F$38,"Descriptive_Churn_Risk_Tier","{tier_name}"),0),{baseline})'
        )
        fmt = workbook.add_format(
            {
                "font_name": "Aptos",
                "font_size": 8,
                "bold": True,
                "font_color": COLORS["navy"],
                "bg_color": COLORS["white"],
                "num_format": f'"{tier_name}  "#,##0',
                "align": "center",
                "valign": "vcenter",
                "border": 1,
                "border_color": COLORS["border"],
            }
        )
        write_formula_merged(dashboard, f"V{row_num}:X{row_num}", formula, fmt, baseline)

    # Executive Insights: named cells provide a future dynamic narrative layer.
    dashboard.merge_range("A40:X40", "EXECUTIVE INSIGHTS", section_band_fmt)
    insight_tiles = [
        ("A41:F42", "CRITICAL FINDING", "=insightCritical", "Month-to-Month customers have a 51.7% churn rate, the highest contract-level exposure.", COLORS["coral"]),
        ("G41:L42", "MAJOR DRIVER", "=insightDriver", "Churn reaches 77.2% among customers in their first 0–6 months.", COLORS["blue"]),
        ("M41:R42", "EARLY WARNING", "=insightWarning", "Month-to-Month + 0–6 months + Fiber records 91.2% churn: 444 of 487 existing customers.", COLORS["coral"]),
        ("S41:X42", "BUSINESS OPPORTUNITY", "=insightOpportunity", "Fiber churn is 42.1%; better devices (313) and better offers (311) lead stated exit reasons.", COLORS["teal"]),
    ]
    for cell_range, label, formula, cached_text, accent in insight_tiles:
        start, end = cell_range.split(":")
        start_col = ord(start[0]) - ord("A")
        start_row = int(start[1:]) - 1
        end_col = ord(end[0]) - ord("A")
        end_row = int(end[1:]) - 1
        add_outer_border(dashboard, start_row, start_col, end_row, end_col, workbook)
        dashboard.merge_range(
            start_row, start_col, start_row, end_col, label,
            workbook.add_format(
                {"font_name": "Aptos", "font_size": 8, "bold": True, "font_color": accent, "bg_color": COLORS["white"], "top": 3, "top_color": accent, "left": 1, "left_color": COLORS["border"], "right": 1, "right_color": COLORS["border"], "align": "left", "valign": "vcenter"}
            ),
        )
        body_range = f"{chr(65 + start_col)}{start_row + 2}:{chr(65 + end_col)}{end_row + 1}"
        body_fmt = workbook.add_format(
            {"font_name": "Aptos", "font_size": 8, "font_color": COLORS["text"], "bg_color": COLORS["white"], "text_wrap": True, "valign": "vcenter", "left": 1, "left_color": COLORS["border"], "right": 1, "right_color": COLORS["border"], "bottom": 1, "bottom_color": COLORS["border"]}
        )
        write_formula_merged(dashboard, body_range, formula, body_fmt, cached_text)
    dashboard.merge_range(
        "A43:X43",
        "Baseline executive context • insight cells are named for future dynamic narratives • associations and stated reasons do not establish causation",
        note_fmt,
    )

    # Management recommendations.
    dashboard.merge_range("A45:X45", "BUSINESS RECOMMENDATIONS", section_band_fmt)
    recommendations = [
        (
            "A46:H53", COLORS["coral"], "01  EARLY-TENURE SAVE JOURNEY",
            "Rationale  |  0–6 month churn is 77.2%; the M2M + 0–6 + Fiber intersection reaches 91.2%.",
            "Action  |  Test onboarding check-ins, service recovery and targeted save treatment during the first 90 days.",
            "EXPECTED IMPACT POTENTIAL  •  HIGH",
        ),
        (
            "I46:P53", COLORS["blue"], "02  COMMITMENT MIGRATION",
            "Rationale  |  Month-to-Month churn is 51.7%, the strongest contract-level exposure.",
            "Action  |  Test transparent one- and two-year migration offers with clear customer value and safeguards.",
            "EXPECTED IMPACT POTENTIAL  •  HIGH",
        ),
        (
            "Q46:X53", COLORS["teal"], "03  COMPETITOR VALUE RESPONSE",
            "Rationale  |  Better devices (313) and better offers (311) lead stated reasons; Fiber churn is 42.1%.",
            "Action  |  Review the device proposition and test competitive offers for relevant Fiber segments.",
            "EXPECTED IMPACT POTENTIAL  •  MEDIUM–HIGH",
        ),
    ]
    for cell_range, accent, title, rationale, action, impact in recommendations:
        start, end = cell_range.split(":")
        start_col = ord(start[0]) - ord("A")
        start_row = int(start[1:]) - 1
        end_col = ord(end[0]) - ord("A")
        end_row = int(end[1:]) - 1
        add_outer_border(dashboard, start_row, start_col, end_row, end_col, workbook)
        dashboard.merge_range(
            start_row, start_col, start_row + 1, end_col, title,
            workbook.add_format(
                {"font_name": "Aptos", "font_size": 10, "bold": True, "font_color": COLORS["navy"], "bg_color": COLORS["white"], "top": 3, "top_color": accent, "left": 1, "left_color": COLORS["border"], "right": 1, "right_color": COLORS["border"], "align": "left", "valign": "vcenter", "text_wrap": True}
            ),
        )
        dashboard.merge_range(
            start_row + 2, start_col, start_row + 3, end_col, rationale,
            workbook.add_format(
                {"font_name": "Aptos", "font_size": 8, "font_color": COLORS["text"], "bg_color": COLORS["white"], "left": 1, "left_color": COLORS["border"], "right": 1, "right_color": COLORS["border"], "text_wrap": True, "valign": "vcenter"}
            ),
        )
        dashboard.merge_range(
            start_row + 4, start_col, start_row + 6, end_col, action,
            workbook.add_format(
                {"font_name": "Aptos", "font_size": 8, "font_color": COLORS["muted"], "bg_color": COLORS["white"], "left": 1, "left_color": COLORS["border"], "right": 1, "right_color": COLORS["border"], "text_wrap": True, "valign": "vcenter"}
            ),
        )
        dashboard.merge_range(
            end_row, start_col, end_row, end_col, impact,
            workbook.add_format(
                {"font_name": "Aptos", "font_size": 8, "bold": True, "font_color": COLORS["white"], "bg_color": accent, "left": 1, "left_color": accent, "right": 1, "right_color": accent, "bottom": 1, "bottom_color": accent, "align": "center", "valign": "vcenter"}
            ),
        )

    # Methodology footer.
    dashboard.merge_range(
        "A55:D55",
        "METHODOLOGY",
        workbook.add_format(
            {"font_name": "Aptos", "font_size": 9, "bold": True, "font_color": COLORS["white"], "bg_color": COLORS["navy"], "align": "left", "valign": "vcenter"}
        ),
    )
    footer_fmt = workbook.add_format(
        {"font_name": "Aptos", "font_size": 8, "font_color": COLORS["muted"], "bg_color": COLORS["white"], "border": 1, "border_color": COLORS["border"], "text_wrap": True, "valign": "vcenter"}
    )
    dashboard.merge_range(
        "A56:R58",
        "Churn Rate = Churned ÷ (Churned + Stayed). Joined customers remain in the source but are excluded from churn and retention denominators. Risk tiers are rule-based descriptive segments, not predictive probabilities. Recommendation impact indicators are directional prioritization signals, not forecasts.",
        footer_fmt,
    )
    dashboard.merge_range(
        "S56:X58",
        "Source\n"
        "data/processed/customers_clean.csv\n"
        "Q2 2022 snapshot",
        workbook.add_format(
            {"font_name": "Aptos", "font_size": 8, "font_color": COLORS["muted"], "bg_color": COLORS["white"], "border": 1, "border_color": COLORS["border"], "align": "right", "valign": "vcenter", "text_wrap": True}
        ),
    )

    # Keep support sheets visible for the required native wiring phase.
    dashboard.activate()
    dashboard.set_selection("A1")

    workbook.close()

    if sha256(SOURCE) != source_hash:
        raise AssertionError("Source CSV changed during workbook build")


def programmatic_qa(
    output: Path, source_hash: str, expected_rows: int = 7043
) -> dict[str, Any]:
    if not output.exists() or output.stat().st_size == 0:
        raise AssertionError("Workbook was not created")
    if sha256(SOURCE) != source_hash:
        raise AssertionError("Source CSV hash changed")

    workbook = load_workbook(output, data_only=False, read_only=False)
    expected_sheets = ["01_Dashboard", "02_Pivots", "03_Calculations", "04_Data"]
    if workbook.sheetnames != expected_sheets:
        raise AssertionError(f"Unexpected sheets: {workbook.sheetnames}")

    data_ws = workbook["04_Data"]
    if "tblCustomers" not in data_ws.tables:
        raise AssertionError("tblCustomers is missing")
    table_ref = data_ws.tables["tblCustomers"].ref
    if table_ref != "A1:BC7044":
        raise AssertionError(f"Unexpected table range: {table_ref}")
    if data_ws.max_row != expected_rows + 1:
        raise AssertionError(f"Unexpected data row count: {data_ws.max_row - 1}")
    if data_ws.max_column != 55:
        raise AssertionError(f"Unexpected data column count: {data_ws.max_column}")

    dashboard = workbook["01_Dashboard"]
    required_merges = {
        "A1:P2",
        "Q1:X2",
        "A9:C10",
        "A13:X13",
        "A14:P14",
        "Q14:X14",
        "A29:M29",
        "N29:X29",
        "A40:X40",
        "A45:X45",
        "A56:R58",
    }
    merge_strings = {str(item) for item in dashboard.merged_cells.ranges}
    missing_merges = required_merges - merge_strings
    if missing_merges:
        raise AssertionError(f"Missing dashboard merges: {sorted(missing_merges)}")

    calc = workbook["03_Calculations"]
    expected_kpi_formulas = {
        "C5": '=IFERROR(GETPIVOTDATA("Existing Customers",\'02_Pivots\'!$A$3),6589)',
        "C6": '=IFERROR(GETPIVOTDATA("Churned Customers",\'02_Pivots\'!$A$3),1869)',
        "C7": '=IFERROR(GETPIVOTDATA("Retained Customers",\'02_Pivots\'!$A$3),4720)',
        "C8": "=IF(C5=0,NA(),C6/C5)",
        "C9": "=IF(C5=0,NA(),C7/C5)",
        "C10": '=IFERROR(GETPIVOTDATA("Monthly Recurring Value Lost",\'02_Pivots\'!$A$3),137086.65)',
    }
    for coordinate, expected_formula in expected_kpi_formulas.items():
        if calc[coordinate].value != expected_formula:
            raise AssertionError(f"Unexpected KPI formula at 03_Calculations!{coordinate}")
    if "nativePivotsReady" not in workbook.defined_names:
        raise AssertionError("nativePivotsReady name is missing")
    expected_names = {
        "nativePivotsReady": "'03_Calculations'!$H$2",
        "kpiExisting": "'03_Calculations'!$C$5",
        "kpiChurned": "'03_Calculations'!$C$6",
        "kpiRetained": "'03_Calculations'!$C$7",
        "kpiChurnRate": "'03_Calculations'!$C$8",
        "kpiRetentionRate": "'03_Calculations'!$C$9",
        "kpiMRVL": "'03_Calculations'!$C$10",
        "insightCritical": "'03_Calculations'!$I$25",
        "insightDriver": "'03_Calculations'!$I$26",
        "insightWarning": "'03_Calculations'!$I$27",
        "insightOpportunity": "'03_Calculations'!$I$28",
    }
    for name, destination in expected_names.items():
        if name not in workbook.defined_names:
            raise AssertionError(f"Missing defined name: {name}")
        if workbook.defined_names[name].attr_text != destination:
            raise AssertionError(f"Unexpected destination for defined name: {name}")
    if calc["I6"].data_type != "s" or not str(calc["I6"].value).startswith("="):
        raise AssertionError("Helper formula documentation must be stored as literal text")

    expected_dashboard_formulas = {
        "A9": "=kpiExisting",
        "E9": "=kpiChurned",
        "I9": "=kpiRetained",
        "M9": "=kpiChurnRate",
        "Q9": "=kpiRetentionRate",
        "U9": "=kpiMRVL",
        "A42": "=insightCritical",
        "G42": "=insightDriver",
        "M42": "=insightWarning",
        "S42": "=insightOpportunity",
    }
    for coordinate, expected_formula in expected_dashboard_formulas.items():
        if dashboard[coordinate].value != expected_formula:
            raise AssertionError(f"Unexpected dashboard formula at {coordinate}")

    for rate_row in range(16, 28, 2):
        for column in ("E", "I", "M"):
            formula = dashboard[f"{column}{rate_row}"].value
            if not isinstance(formula, str) or not formula.startswith(
                '=IF(nativePivotsReady,IFERROR(GETPIVOTDATA("Churn Rate"'
            ):
                raise AssertionError(
                    f"Heatmap rate formula is missing at {column}{rate_row}"
                )
    for row in range(32, 36):
        formula = dashboard[f"V{row}"].value
        if not isinstance(formula, str) or "GETPIVOTDATA" not in formula:
            raise AssertionError(f"Risk volume formula is missing at V{row}")

    formula_errors: list[str] = []
    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and ("#REF!" in value or "#DIV/0!" in value):
                    formula_errors.append(f"{ws.title}!{cell.coordinate}: {value}")
    if formula_errors:
        raise AssertionError(f"Detected broken references: {formula_errors[:5]}")

    with zipfile.ZipFile(output) as archive:
        names = archive.namelist()
        if any(name.startswith("xl/externalLinks/") for name in names):
            raise AssertionError("External links detected")
        if "xl/tables/table1.xml" not in names:
            raise AssertionError("Excel Table XML is missing")
        sheet3_xml = archive.read("xl/worksheets/sheet3.xml")
        if b"<f>--OR([@[" in sheet3_xml or b"<f>IF([@" in sheet3_xml:
            raise AssertionError(
                "Invalid table-row documentation formula was serialized in sheet3.xml"
            )

    style_checks = {
        "header_fill": dashboard["A1"].fill.fgColor.rgb,
        "kpi_number_format": dashboard["A9"].number_format,
        "heatmap_number_format": dashboard["E16"].number_format,
        "dashboard_zoom": dashboard.sheet_view.zoomScale,
        "gridlines_hidden": dashboard.sheet_view.showGridLines is False,
        "headings_hidden": dashboard.sheet_view.showRowColHeaders is False,
    }
    if dashboard.sheet_view.zoomScale != 75:
        raise AssertionError("Dashboard zoom is not 75%")
    if dashboard.sheet_view.showGridLines is not False:
        raise AssertionError("Dashboard gridlines are visible")
    if dashboard.sheet_view.showRowColHeaders is not False:
        raise AssertionError("Dashboard headings are visible")
    if dashboard.max_row < 58 or dashboard.max_column < 24:
        raise AssertionError("Dashboard canvas is incomplete")

    workbook.close()
    return {
        "workbook": str(output),
        "size_bytes": output.stat().st_size,
        "sheets": expected_sheets,
        "table": "tblCustomers",
        "table_ref": table_ref,
        "data_rows": expected_rows,
        "data_columns": 55,
        "style_checks": style_checks,
        "source_sha256": source_hash,
    }


def print_validation(validation: dict[str, Any], qa: dict[str, Any]) -> None:
    metrics = validation["metrics"]
    print("=== SOURCE VALIDATION ===")
    print(f"[PASS] Total Customers: {metrics['total']:,}")
    print(f"[PASS] Existing Customers: {metrics['existing']:,}")
    print(f"[PASS] Churned Customers: {metrics['churned']:,}")
    print(f"[PASS] Retained Customers: {metrics['retained']:,}")
    print(f"[PASS] Joined Customers: {metrics['joined']:,}")
    print(f"[PASS] Churn Rate: {metrics['churn_rate']:.2%}")
    print(f"[PASS] Retention Rate: {metrics['retention_rate']:.2%}")
    print(f"[PASS] MRVL: ${metrics['mrvl']:,.2f}")
    print("=== WORKBOOK QA ===")
    print(f"[PASS] Workbook: {qa['workbook']}")
    print(f"[PASS] Sheets: {', '.join(qa['sheets'])}")
    print(f"[PASS] Table: {qa['table']} ({qa['table_ref']})")
    print(f"[PASS] Data: {qa['data_rows']:,} rows × {qa['data_columns']} columns")
    print(f"[PASS] Size: {qa['size_bytes'] / 1024 / 1024:.2f} MB")
    print("[PASS] Formulas, names, styles, merges, view settings, and links validated")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Workbook output path (default: {DEFAULT_OUTPUT})",
    )
    args = parser.parse_args()

    source_hash = sha256(SOURCE)
    df, validation = load_validate_and_enrich()
    build_workbook(args.output.resolve(), df, validation, source_hash)
    qa = programmatic_qa(args.output.resolve(), source_hash)
    print_validation(validation, qa)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # concise build failure for CLI use
        print(f"[FAIL] {type(exc).__name__}: {exc}", file=sys.stderr)
        sys.exit(1)
